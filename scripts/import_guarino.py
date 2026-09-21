import json
import re
import sys
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen
import ssl
import certifi

from openpyxl import load_workbook


# ============================================================
# CONFIGURACIÓN
# ============================================================

API_URL = "https://backend-production-c8a3.up.railway.app"

EXCEL_FILE = Path(__file__).parent / "Guarino_Gabriel_20-33186724-2.xlsx"

CLIENT_NAME = "Gabriel Guarino"
CLIENT_CUIT = "20-33186724-2"

# IMPORTANTE:
# Primero ejecutamos siempre en True.
# Cuando veamos que el análisis es correcto, lo cambiamos a False.
DRY_RUN = False


# ============================================================
# UTILIDADES
# ============================================================

def excel_date(value):
    """Convierte número serial de Excel a YYYY-MM-DD."""
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, (int, float)):
        date = datetime(1899, 12, 30) + timedelta(days=value)
        return date.date().isoformat()

    return None


def decimal_value(value):
    if value is None:
        return None

    return Decimal(str(value))


def api_request(method, endpoint, data=None):
    url = f"{API_URL.rstrip('/')}/{endpoint.lstrip('/')}"

    body = None

    if data is not None:
        body = json.dumps(data).encode("utf-8")

    request = Request(
        url,
        data=body,
        method=method,
        headers={
            "Content-Type": "application/json",
        },
    )

    try:
        ssl_context = ssl.create_default_context(
            cafile=certifi.where()
        )

        with urlopen(request, context=ssl_context) as response:
            content = response.read().decode("utf-8")

            if not content:
                return None

            return json.loads(content)

    except HTTPError as exc:
        content = exc.read().decode("utf-8")

        raise RuntimeError(
            f"{method} {url} -> {exc.code}: {content}"
        ) from exc


def extract_ticker(tipo_mov):
    """
    Ejemplos:
        Compra(KO) -> KO
        Compra(GGAL) GRUPO... -> GGAL
        Venta(NVDA) -> NVDA
        Transferencia de Titulos IN - (WMT) -> WMT
    """
    if not tipo_mov:
        return None

    match = re.search(r"\(([^)]+)\)", str(tipo_mov))

    if match:
        return match.group(1).strip().upper()

    return None


def extract_instrument_name(tipo_mov, ticker):
    if not tipo_mov or not ticker:
        return ticker

    text = str(tipo_mov)

    # Elimina la parte anterior hasta "(TICKER)"
    marker = f"({ticker})"
    position = text.upper().find(marker.upper())

    if position == -1:
        return ticker

    name = text[position + len(marker):].strip()

    if not name:
        return ticker

    return name


def total_fees(row):
    """
    Comisión + IVA comisión + otros impuestos.
    Lo conservamos para tener el dato disponible.
    """
    values = [
        row.get("Comis."),
        row.get("Iva Com."),
        row.get("Otros Imp."),
    ]

    total = Decimal("0")

    for value in values:
        if value is not None:
            total += Decimal(str(value))

    return total


# ============================================================
# LECTURA DEL EXCEL
# ============================================================

def read_excel():
    workbook = load_workbook(
        EXCEL_FILE,
        data_only=True,
    )

    worksheet = workbook.active

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    rows = []

    for excel_row_number, values in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        row = dict(zip(headers, values))
        row["_excel_row"] = excel_row_number
        rows.append(row)

    return rows


# ============================================================
# CLASIFICACIÓN
# ============================================================

def classify_rows(rows):
    operations = []
    ignored = []
    review = []

    current_ticker = None
    current_name = None

    for row in rows:
        tipo_mov = row.get("Tipo Mov.")
        explicit_ticker = extract_ticker(tipo_mov)

        # El Excel está organizado por bloques de instrumento.
        if explicit_ticker:
            current_ticker = explicit_ticker
            current_name = extract_instrument_name(
                tipo_mov,
                explicit_ticker,
            )

        ticker = explicit_ticker or current_ticker

        quantity = decimal_value(
            row.get("Cant. titulos")
        )

        operation_date = excel_date(
            row.get("Concert.")
        )

        liquid = row.get("Liquid.")

        # ----------------------------------------------------
        # FILAS VACÍAS
        # ----------------------------------------------------

        if (
            tipo_mov is None
            and quantity is None
            and row.get("costo") is None
            and row.get("Monto") is None
        ):
            ignored.append(
                (row, "Fila vacía")
            )
            continue

        tipo_text = (
            str(tipo_mov).strip()
            if tipo_mov
            else ""
        )

        # ----------------------------------------------------
        # DIVIDENDOS
        # ----------------------------------------------------

        if tipo_text.startswith("Pago de Dividendos"):
            ignored.append(
                (row, "Dividendo - todavía no soportado")
            )
            continue

        # ----------------------------------------------------
        # FCI
        # ----------------------------------------------------

        if tipo_text.startswith("Suscripción FCI"):
            review.append(
                (row, "Suscripción FCI - fuera del MVP actual")
            )
            continue

        # ----------------------------------------------------
        # TRANSFERENCIAS
        # ----------------------------------------------------

        if tipo_text.startswith("Transferencia de Titulos IN"):
            if quantity is None or quantity <= 0:
                review.append(
                    (row, "Split / transferencia sin cantidad válida")
                )
                continue

            if ticker is None or operation_date is None:
                review.append(
                    (row, "Split / transferencia incompleta")
                )
                continue

            operations.append(
                {
                    "excel_row": row["_excel_row"],
                    "type": "BUY",
                    "ticker": ticker,
                    "name": current_name or ticker,
                    "date": operation_date,
                    "quantity": quantity,
                    "unit_price": Decimal("0"),
                    "commission": Decimal("0"),
                    "source_fees": Decimal("0"),
                    "reason": "Split - títulos incorporados a costo 0",
                }
            )

            continue

        # ----------------------------------------------------
        # POSICIÓN INICIAL
        # ----------------------------------------------------

        if str(liquid).strip().upper() == "E INICIAL":
            cost = decimal_value(
                row.get("costo")
            )

            if (
                    ticker is None
                    or quantity is None
                    or quantity <= 0
                    or cost is None
                    or operation_date is None
            ):
                review.append(
                    (row, "Posición inicial incompleta")
                )
                continue

            operations.append(
                {
                    "excel_row": row["_excel_row"],
                    "type": "BUY",
                    "ticker": ticker,
                    "name": current_name or ticker,
                    "date": operation_date,
                    "quantity": quantity,
                    "unit_price": cost,
                    "commission": Decimal("0"),
                    "source_fees": Decimal("0"),
                    "reason": "Posición inicial",
                }
            )

            continue

        # ----------------------------------------------------
        # COMPRA NORMAL
        # ----------------------------------------------------

        if tipo_text.startswith("Compra("):
            cost = decimal_value(
                row.get("costo")
            )

            if (
                    ticker is None
                    or quantity is None
                    or quantity <= 0
                    or operation_date is None
                    or cost is None
            ):
                review.append(
                    (row, "Compra incompleta")
                )
                continue

            operations.append(
                {
                    "excel_row": row["_excel_row"],
                    "type": "BUY",
                    "ticker": ticker,
                    "name": current_name or ticker,
                    "date": operation_date,
                    "quantity": quantity,
                    "unit_price": cost,

                    # El costo histórico ya incorpora los gastos.
                    "commission": Decimal("0"),

                    "source_fees": total_fees(row),
                    "reason": "Compra normal",
                }
            )

            continue

        # ----------------------------------------------------
        # VENTA
        # ----------------------------------------------------

        if (
                tipo_text.upper() == "VENTA"
                or tipo_text.startswith("Venta(")
        ):
            if ticker is None:
                review.append(
                    (row, "Venta sin instrumento identificable")
                )
                continue

            if quantity is None or quantity == 0:
                review.append(
                    (row, "Venta sin cantidad")
                )
                continue

            if operation_date is None:
                review.append(
                    (row, "Venta sin fecha")
                )
                continue

            sale_quantity = abs(quantity)

            # Primero intentamos usar el precio explícito
            sale_price = decimal_value(
                row.get("Precio")
            )

            reason = "Venta"

            # Si la venta no tiene Precio explícito,
            # utilizamos el costo unitario informado por el Excel.
            #
            # Ejemplo YPFD:
            # costo = 21020.932222...
            if sale_price is None:
                sale_cost = decimal_value(
                    row.get("costo")
                )

                if sale_cost is None:
                    review.append(
                        (
                            row,
                            f"Venta de {ticker} sin precio ni costo",
                        )
                    )
                    continue

                sale_price = abs(sale_cost)

                reason = (
                    "Venta - precio unitario tomado de costo"
                )

            operations.append(
                {
                    "excel_row": row["_excel_row"],
                    "type": "SELL",
                    "ticker": ticker,
                    "name": current_name or ticker,
                    "date": operation_date,
                    "quantity": sale_quantity,
                    "unit_price": sale_price,
                    "commission": total_fees(row),
                    "source_fees": total_fees(row),
                    "reason": reason,
                }
            )

            continue
        # ----------------------------------------------------
        # SUBTOTALES / VALUACIONES
        # ----------------------------------------------------

        if tipo_mov is None:
            ignored.append(
                (
                    row,
                    "Subtotal / valuación / fila derivada",
                )
            )
            continue

        # ----------------------------------------------------
        # DESCONOCIDO
        # ----------------------------------------------------

        review.append(
            (
                row,
                f"Tipo de movimiento no reconocido: {tipo_text}",
            )
        )

    return operations, ignored, review


# ============================================================
# API FIFOLIO
# ============================================================

def get_or_create_client():
    clients = api_request(
        "GET",
        "/clients/",
    )

    for client in clients:
        if client.get("cuit") == CLIENT_CUIT:
            return client

    return api_request(
        "POST",
        "/clients/",
        {
            "name": CLIENT_NAME,
            "cuit": CLIENT_CUIT,
        },
    )


def get_or_create_instrument(ticker, name, currency):
    instruments = api_request(
        "GET",
        "/instruments/",
    )

    for instrument in instruments:
        if instrument["ticker"].upper() == ticker.upper():
            return instrument

    return api_request(
        "POST",
        "/instruments/",
        {
            "ticker": ticker,
            "name": name or ticker,
            "currency": currency,
        },
    )

def ensure_client_has_no_operations(client):
    """
    Protección para evitar importar dos veces la cartera.

    Este importador es de carga inicial: sólo puede ejecutarse
    si el cliente todavía no tiene operaciones en FIFolio.
    """
    existing_operations = api_request(
        "GET",
        "/operations/",
    )

    client_operations = [
        operation
        for operation in existing_operations
        if operation.get("client_id") == client["id"]
    ]

    if client_operations:
        raise RuntimeError(
            "\nIMPORTACIÓN CANCELADA.\n"
            f'El cliente {client["name"]} '
            f'({client["cuit"]}) ya tiene '
            f"{len(client_operations)} operación(es) en FIFolio.\n"
            "No se importará nada para evitar duplicados."
        )
def import_operations(operations):
    client = get_or_create_client()

    # Protección contra doble importación.
    ensure_client_has_no_operations(client)

    instruments = {}

    # Primero creamos todos los instrumentos.
    for operation in operations:
        ticker = operation["ticker"]

        if ticker in instruments:
            continue

        # El Excel de Guarino tiene los títulos negociados
        # principalmente en pesos.
        currency = "ARS"

        instruments[ticker] = get_or_create_instrument(
            ticker=ticker,
            name=operation["name"],
            currency=currency,
        )

    # CRÍTICO:
    # FIFO depende del orden cronológico.
    operations.sort(
        key=lambda x: (
            x["date"],
            x["excel_row"],
        )
    )

    created = []

    for operation in operations:
        instrument = instruments[
            operation["ticker"]
        ]

        payload = {
            "client_id": client["id"],
            "instrument_id": instrument["id"],
            "operation_date": operation["date"],
            "quantity": str(operation["quantity"]),
            "unit_price": str(operation["unit_price"]),
            "commission": str(operation["commission"]),
        }

        endpoint = (
            "/operations/buy"
            if operation["type"] == "BUY"
            else "/operations/sell"
        )

        result = api_request(
            "POST",
            endpoint,
            payload,
        )

        created.append(
            {
                **operation,
                "fifolio_id": result["id"],
            }
        )

        print(
            f'OK fila {operation["excel_row"]}: '
            f'{operation["type"]} '
            f'{operation["quantity"]} '
            f'{operation["ticker"]} '
            f'@ {operation["unit_price"]}'
        )

    return created


# ============================================================
# REPORTE
# ============================================================

def print_report(operations, ignored, review):
    buys = [
        op for op in operations
        if op["type"] == "BUY"
    ]

    sells = [
        op for op in operations
        if op["type"] == "SELL"
    ]

    initials = [
        op for op in operations
        if op["reason"] == "Posición inicial"
    ]

    tickers = sorted(
        set(op["ticker"] for op in operations)
    )

    print()
    print("=" * 70)
    print("ANÁLISIS IMPORTACIÓN GUARINO")
    print("=" * 70)

    print(f"Operaciones importables : {len(operations)}")
    print(f"Compras                 : {len(buys)}")
    print(f"  Posiciones iniciales  : {len(initials)}")
    print(f"Ventas                  : {len(sells)}")
    print(f"Ignoradas               : {len(ignored)}")
    print(f"Requieren revisión      : {len(review)}")

    print()
    print("Instrumentos:")
    print(", ".join(tickers))

    print()
    print("-" * 70)
    print("OPERACIONES IMPORTABLES")
    print("-" * 70)

    for op in sorted(
        operations,
        key=lambda x: (
            x["date"],
            x["excel_row"],
        ),
    ):
        print(
            f'Fila {op["excel_row"]:>3} | '
            f'{op["date"]} | '
            f'{op["type"]:4} | '
            f'{op["ticker"]:8} | '
            f'{op["quantity"]} @ {op["unit_price"]} | '
            f'{op["reason"]}'
        )

    print()
    print("-" * 70)
    print("REQUIEREN REVISIÓN")
    print("-" * 70)

    for row, reason in review:
        print(
            f'Fila {row["_excel_row"]:>3}: {reason}'
        )

    print()
    print("-" * 70)

    if DRY_RUN:
        print("DRY_RUN=True")
        print("NO SE MODIFICÓ FIFOLIO.")
    else:
        print("DRY_RUN=False")
        print("SE IMPORTARÁN DATOS A FIFOLIO.")

    print("=" * 70)


EXPECTED_POSITIONS = {
    "YPFD": Decimal("0"),
    "KO": Decimal("47"),
    "MCD": Decimal("46"),
    "WMT": Decimal("72"),
    "GOOGL": Decimal("323"),
    "AAPL": Decimal("138"),
    "NVDA": Decimal("91"),
    "SPY": Decimal("31"),
    "QQQ": Decimal("10"),
    "LECAO": Decimal("1960"),
    "GGAL": Decimal("48"),
    "CVX": Decimal("4"),
    "GE": Decimal("24"),
    "PFE": Decimal("6"),
    "T": Decimal("8"),
    "VALE": Decimal("4"),
    "VIST": Decimal("16"),
    "BRKB": Decimal("27"),
    "MO": Decimal("2"),
}


def print_reconciliation(operations):
    positions = {}

    for op in operations:
        ticker = op["ticker"]

        if ticker not in positions:
            positions[ticker] = Decimal("0")

        if op["type"] == "BUY":
            positions[ticker] += op["quantity"]

        elif op["type"] == "SELL":
            positions[ticker] -= op["quantity"]

    print()
    print("=" * 82)
    print("CONCILIACIÓN DE POSICIONES")
    print("=" * 82)

    print(
        f'{"TICKER":<10}'
        f'{"FIFOLIO":>15}'
        f'{"EXCEL":>15}'
        f'{"DIFERENCIA":>15}'
        f'{"ESTADO":>15}'
    )

    print("-" * 82)

    all_tickers = sorted(
        set(positions)
        | set(EXPECTED_POSITIONS)
    )

    differences = []

    for ticker in all_tickers:
        fifolio = positions.get(
            ticker,
            Decimal("0"),
        )

        excel = EXPECTED_POSITIONS.get(
            ticker
        )

        if excel is None:
            print(
                f"{ticker:<10}"
                f"{str(fifolio):>15}"
                f"{'?':>15}"
                f"{'?':>15}"
                f"{'REVISAR':>15}"
            )
            continue

        difference = fifolio - excel

        if difference == 0:
            status = "OK"
        else:
            status = "PENDIENTE"
            differences.append(
                (ticker, fifolio, excel, difference)
            )

        print(
            f"{ticker:<10}"
            f"{str(fifolio):>15}"
            f"{str(excel):>15}"
            f"{str(difference):>15}"
            f"{status:>15}"
        )

    print("=" * 82)

    if not differences:
        print("CONCILIACIÓN PERFECTA.")
        print(
            "Todas las posiciones coinciden con el Excel."
        )

    else:
        print(
            f"Hay {len(differences)} instrumento(s) "
            "pendientes de conciliación:"
        )

        for ticker, fifolio, excel, difference in differences:
            print(
                f"  {ticker}: "
                f"FIFolio={fifolio}, "
                f"Excel={excel}, "
                f"diferencia={difference}"
            )

    print("=" * 82)

# ============================================================
# MAIN
# ============================================================

def main():
    if not EXCEL_FILE.exists():
        print(
            f"No se encontró el Excel:\n{EXCEL_FILE}"
        )
        sys.exit(1)

    rows = read_excel()

    operations, ignored, review = classify_rows(
        rows
    )

    print_report(
        operations,
        ignored,
        review,
    )

    print_reconciliation(operations)

    if DRY_RUN:
        return

    confirmation = input(
        "\nEscribí IMPORTAR para continuar: "
    )

    if confirmation != "IMPORTAR":
        print("Importación cancelada.")
        return

    created = import_operations(
        operations
    )

    print()
    print(
        f"Importación terminada: "
        f"{len(created)} operaciones creadas."
    )


if __name__ == "__main__":
    main()